import json
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from google.genai import errors

from app.contracts import ComparedField
from app.extraction import (
    ExtractionFailure,
    ExtractionFailureCode,
    GeminiDocument,
    GeminiExtractor,
    grounded_extraction,
    local_extraction,
    scan_extraction,
)
from app.formats import parse_document, preflight
from app.gemini import GeminiCallError, GeminiNotConfigured, KeyAttempt

ATTACHMENTS = (
    Path(__file__).resolve().parents[3]
    / "data"
    / "sdoc-hackathon-bundle"
    / "attachments"
)
OK = (KeyAttempt(key_index=1, outcome="SUCCEEDED", status_code=None),)


def _field(value, page=None, region=None):
    return {"value": value, "page": page, "region": region}


def _answer(**overrides):
    fields = {
        "shipper": _field("ACME LTD", 1, "party"),
        "consignee": _field("BETA LTD", 1, "party"),
        "notify_party": _field("BETA LTD", 1, "party"),
        "port_of_loading": _field("SINGAPORE", 1, "routing"),
        "port_of_discharge": _field("BUSAN, SOUTH KOREA", 1, "routing"),
        "container_count": _field("2 x 40'HC", 1, "cargo"),
        "gross_weight_kg": _field("40,326 KG", 1, "cargo"),
    }
    fields.update(overrides)
    return json.dumps(
        {
            "document_title": "BILL OF LADING (DRAFT)",
            "transcription": "BILL OF LADING",
            "fields": fields,
        }
    )


def _generate(text=None, error=None, key_attempts=OK, calls=None):
    async def generate(contents, config=None, attempts=None):
        if calls is not None:
            calls.append((contents, config))
        if error is not None:
            raise error
        return SimpleNamespace(
            text=text, model_version="gemini-3.5-flash-001"
        ), key_attempts

    return generate


@pytest.mark.asyncio
async def test_scan_is_sent_as_pdf_bytes_with_a_json_schema():
    calls = []
    extractor = GeminiExtractor(_generate(_answer(), calls=calls))

    outcome = await extractor.read_scan(b"%PDF-1.5 scan")

    contents, config = calls[0]
    assert contents[0].inline_data.mime_type == "application/pdf"
    assert config.response_mime_type == "application/json"
    assert config.temperature == 0
    fields_schema = config.response_json_schema["$defs"]["GeminiFields"]["properties"]
    assert set(fields_schema) == {
        "shipper",
        "consignee",
        "notify_party",
        "port_of_loading",
        "port_of_discharge",
        "container_count",
        "gross_weight_kg",
    }
    assert outcome.document.fields.shipper.value == "ACME LTD"
    assert outcome.model_version == "gemini-3.5-flash-001"


def test_scan_values_get_approximate_page_and_region_anchors():
    outcome = SimpleNamespace(
        document=GeminiDocument.model_validate_json(_answer()),
        key_attempts=OK,
    )
    result = scan_extraction(
        outcome, attachment_id="a", file_name="s.pdf", page_count=1
    )

    provenance = result.values[0].provenance.root
    assert provenance.format == "scanned_pdf"
    assert provenance.location.model_dump() == {
        "kind": "scanned_pdf",
        "page": 1,
        "approximate": True,
        "region": "party",
    }


@pytest.mark.parametrize(
    "bad_field",
    [
        _field("ACME LTD", None, "party"),
        _field("ACME LTD", 2, "party"),
        _field("ACME LTD", 1, None),
    ],
)
def test_scan_value_without_a_valid_page_and_region_fails_closed(bad_field):
    outcome = SimpleNamespace(
        document=GeminiDocument.model_validate_json(_answer(shipper=bad_field)),
        key_attempts=OK,
    )
    with pytest.raises(ExtractionFailure) as caught:
        scan_extraction(outcome, attachment_id="a", file_name="s.pdf", page_count=1)
    assert caught.value.code is ExtractionFailureCode.INVALID_SCHEMA


def test_absent_scan_value_is_omitted_not_fabricated():
    outcome = SimpleNamespace(
        document=GeminiDocument.model_validate_json(
            _answer(gross_weight_kg=_field(None))
        ),
        key_attempts=OK,
    )
    result = scan_extraction(
        outcome, attachment_id="a", file_name="s.pdf", page_count=1
    )

    assert ComparedField.GROSS_WEIGHT_KG not in {value.field for value in result.values}


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "text",
    [
        "not json",
        json.dumps({"document_title": "x", "transcription": "x"}),
        _answer(shipper={"value": 3}),
    ],
)
async def test_invalid_structured_answer_is_invalid_schema(text):
    with pytest.raises(ExtractionFailure) as caught:
        await GeminiExtractor(_generate(text)).read_scan(b"%PDF-")
    assert caught.value.code is ExtractionFailureCode.INVALID_SCHEMA
    assert caught.value.retryable is True
    assert caught.value.key_attempts == OK


def _quota_error(quota_id: str) -> errors.ClientError:
    """A Gemini 429 body; every quota's message says "exceeded your quota"."""
    return errors.ClientError(
        429,
        {
            "error": {
                "code": 429,
                "message": "You exceeded your current quota, please check your "
                "plan and billing details.",
                "status": "RESOURCE_EXHAUSTED",
                "details": [
                    {
                        "@type": "type.googleapis.com/google.rpc.QuotaFailure",
                        "violations": [
                            {
                                "quotaMetric": "generativelanguage.googleapis.com/"
                                "generate_content_free_tier_requests",
                                "quotaId": quota_id,
                            }
                        ],
                    }
                ],
            }
        },
    )


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("error", "code", "retryable"),
    [
        (
            _quota_error("GenerateRequestsPerMinutePerProjectPerModel-FreeTier"),
            ExtractionFailureCode.RATE_LIMITED,
            True,
        ),
        (
            _quota_error("GenerateRequestsPerDayPerProjectPerModel-FreeTier"),
            ExtractionFailureCode.QUOTA_EXHAUSTED,
            True,
        ),
        (
            errors.ClientError(
                429, {"error": {"message": "Quota of 50 requests per day exceeded"}}
            ),
            ExtractionFailureCode.QUOTA_EXHAUSTED,
            True,
        ),
        (
            errors.ClientError(429, {"error": {"message": "Too many requests"}}),
            ExtractionFailureCode.RATE_LIMITED,
            True,
        ),
        (
            errors.ClientError(400, {"error": {"message": "bad"}}),
            ExtractionFailureCode.PROVIDER_REJECTED,
            False,
        ),
        (
            errors.ServerError(503, {"error": {"message": "unavailable"}}),
            ExtractionFailureCode.PROVIDER_ERROR,
            True,
        ),
    ],
)
async def test_provider_errors_fail_closed_with_distinct_codes(error, code, retryable):
    attempts = (
        KeyAttempt(key_index=1, outcome="RATE_LIMITED", status_code=429),
        KeyAttempt(key_index=2, outcome="FAILED", status_code=error.code),
    )
    extractor = GeminiExtractor(_generate(error=GeminiCallError(error, attempts)))

    with pytest.raises(ExtractionFailure) as caught:
        await extractor.read_scan(b"%PDF-")

    assert caught.value.code is code
    assert caught.value.retryable is retryable
    assert caught.value.key_attempts == attempts


@pytest.mark.asyncio
async def test_slow_gemini_call_times_out():
    import asyncio

    async def slow(contents, config=None, attempts=None):
        await asyncio.sleep(1)

    with pytest.raises(ExtractionFailure) as caught:
        await GeminiExtractor(slow, timeout_seconds=0.01).read_scan(b"%PDF-")
    assert caught.value.code is ExtractionFailureCode.TIMEOUT


@pytest.mark.asyncio
async def test_timed_out_call_still_reports_completed_key_attempts():
    import asyncio

    async def rate_limited_then_hangs(contents, config=None, attempts=None):
        if attempts is not None:
            attempts.append(
                KeyAttempt(key_index=1, outcome="RATE_LIMITED", status_code=429)
            )
        await asyncio.sleep(1)

    with pytest.raises(ExtractionFailure) as caught:
        await GeminiExtractor(rate_limited_then_hangs, timeout_seconds=0.01).read_scan(
            b"%PDF-"
        )
    assert caught.value.code is ExtractionFailureCode.TIMEOUT
    assert caught.value.key_attempts == (
        KeyAttempt(key_index=1, outcome="RATE_LIMITED", status_code=429),
    )


@pytest.mark.asyncio
async def test_missing_key_is_unconfigured():
    with pytest.raises(ExtractionFailure) as caught:
        await GeminiExtractor(_generate(error=GeminiNotConfigured("no key"))).read_scan(
            b"%PDF-"
        )
    assert caught.value.code is ExtractionFailureCode.UNCONFIGURED
    assert caught.value.retryable is False


def _parsed(text: bytes):
    return parse_document(
        text, preflight(text, file_name="x.txt"), attachment_id="a", file_name="x.txt"
    )


def test_local_extraction_keeps_local_anchors():
    data = (ATTACHMENTS / "email_001_SI.txt").read_bytes()
    parsed = parse_document(
        data, preflight(data, file_name="si.txt"), attachment_id="a", file_name="si.txt"
    )

    result = local_extraction(parsed)

    assert [value.field for value in result.values] == list(ComparedField)
    assert all(value.provenance.root.format == "txt" for value in result.values)


def test_grounded_extraction_anchors_model_values_in_the_original_format():
    parsed = _parsed(
        b"SHIPPING INSTRUCTION\nShipper Name: ACME LTD\nConsignee: BETA LTD\n"
        b"Notify: BETA LTD\nPOL: SINGAPORE\nPOD: BUSAN, SOUTH KOREA\n"
        b"Container Count: 2 x 40'HC\nGross Weight (KG): 40,326 KG\n"
    )
    assert parsed.ambiguous_fields == ()  # "Shipper Name" still starts with "shipper"
    parsed = _parsed(b"SI\nSender: ACME LTD\nConsignee: BETA LTD\n")
    absent = {
        name: _field(None)
        for name in (
            "notify_party",
            "port_of_loading",
            "port_of_discharge",
            "container_count",
            "gross_weight_kg",
        )
    }
    outcome = SimpleNamespace(
        document=GeminiDocument.model_validate_json(
            _answer(shipper=_field("ACME LTD"), **absent)
        ),
        key_attempts=OK,
    )

    result = grounded_extraction(parsed, outcome)
    shipper = next(
        value for value in result.values if value.field is ComparedField.SHIPPER
    )

    assert shipper.provenance.root.format == "txt"
    assert shipper.provenance.root.location.line == 2
    consignee = next(
        value for value in result.values if value.field is ComparedField.CONSIGNEE
    )
    assert consignee.raw_value == "BETA LTD"


def test_ungrounded_model_value_fails_closed():
    parsed = _parsed(b"SI\nSender: ACME LTD\n")
    outcome = SimpleNamespace(
        document=GeminiDocument.model_validate_json(
            _answer(shipper=_field("INVENTED LTD"))
        ),
        key_attempts=OK,
    )
    with pytest.raises(ExtractionFailure) as caught:
        grounded_extraction(parsed, outcome)
    assert caught.value.code is ExtractionFailureCode.UNGROUNDED_VALUE


# Each builder prints the shipper's name first on an unlabelled remarks line,
# then shipper, consignee, and container count under their own labels, and
# says where each of those four values is printed.
PDF_BASELINES = (72, 100, 128, 156, 184)


def _txt_labelled():
    data = (
        b"Remarks: ACME LTD\nShipper: ACME LTD\nConsignee: BETA LTD\n"
        b"No. of Containers: 10 x 40'HC\n"
    )
    return data, {
        "remarks": {"kind": "txt", "line": 1, "start_col": 9, "end_col": 17},
        "shipper": {"kind": "txt", "line": 2, "start_col": 9, "end_col": 17},
        "consignee": {"kind": "txt", "line": 3, "start_col": 11, "end_col": 19},
        "containers": {"kind": "txt", "line": 4, "start_col": 19, "end_col": 29},
    }


def _xlsx_labelled():
    import openpyxl

    workbook = openpyxl.Workbook()
    for row in (
        ("Remarks", "ACME LTD"),
        ("Shipper", "ACME LTD"),
        ("Consignee", "BETA LTD"),
        ("No. of Containers", "10 x 40'HC"),
    ):
        workbook.active.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    cells = {"remarks": "B1", "shipper": "B2", "consignee": "B3", "containers": "B4"}
    return buffer.getvalue(), {
        name: {"kind": "xlsx", "sheet": "Sheet", "cell": cell}
        for name, cell in cells.items()
    }


def _docx_labelled():
    import docx

    source = docx.Document()
    source.add_paragraph("Remarks: ACME LTD")
    source.add_paragraph("Consignee: BETA LTD")
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Shipper", "ACME LTD"
    table.cell(1, 0).text, table.cell(1, 1).text = "No. of Containers", "10 x 40'HC"
    buffer = BytesIO()
    source.save(buffer)
    return buffer.getvalue(), {
        "remarks": {"kind": "docx_paragraph", "paragraph_index": 0},
        "consignee": {"kind": "docx_paragraph", "paragraph_index": 1},
        "shipper": {
            "kind": "docx_table",
            "table_index": 0,
            "row_index": 0,
            "col_index": 1,
        },
        "containers": {
            "kind": "docx_table",
            "table_index": 0,
            "row_index": 1,
            "col_index": 1,
        },
    }


def _pdf_labelled():
    import pymupdf

    lines = (
        "Remarks: ACME LTD",
        "Shipper: ACME LTD",
        "Consignee",
        "BETA LTD",
        "No. of Containers: 10 x 40'HC",
    )
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        for baseline, text in zip(PDF_BASELINES, lines):
            page.insert_text((72, baseline), text)
        data = pdf.tobytes()
    # A digital-PDF spot is the baseline of the printed line.
    return data, {"remarks": 72, "shipper": 100, "consignee": 156, "containers": 184}


def _labelled(build):
    data, spots = build()
    document = parse_document(
        data,
        preflight(data, file_name="labelled"),
        attachment_id="a",
        file_name="labelled",
    )
    return document, spots


def _spot(provenance):
    """Where a provenance points; for a digital PDF, the baseline of its line."""
    if provenance is None:
        return None
    location = provenance.root.location.model_dump()
    if location["kind"] != "digital_pdf":
        return location
    _, y0, _, y1 = location["bbox"]
    return next(baseline for baseline in PDF_BASELINES if y0 < baseline < y1)


EVERY_FORMAT = pytest.mark.parametrize(
    "build",
    [_txt_labelled, _xlsx_labelled, _docx_labelled, _pdf_labelled],
    ids=["txt", "xlsx", "docx", "pdf"],
)


@EVERY_FORMAT
def test_locate_matches_whole_tokens_only(build):
    document, spots = _labelled(build)

    assert document.locate("40") is None  # printed only inside 40'HC
    assert document.locate("1") is None  # printed only inside 10
    assert document.locate("ACM") is None  # printed only inside ACME
    assert (
        _spot(document.locate("10 x 40'HC", ComparedField.CONTAINER_COUNT))
        == spots["containers"]
    )


@EVERY_FORMAT
def test_locate_prefers_text_under_the_fields_own_label(build):
    document, spots = _labelled(build)

    assert _spot(document.locate("ACME LTD")) == spots["remarks"]
    assert _spot(document.locate("ACME LTD", ComparedField.SHIPPER)) == spots["shipper"]
    # A field printed under no label of its own falls back to unlabelled text.
    assert (
        _spot(document.locate("ACME LTD", ComparedField.NOTIFY_PARTY))
        == spots["remarks"]
    )


@EVERY_FORMAT
def test_value_printed_only_under_another_fields_label_is_ungrounded(build):
    document, spots = _labelled(build)
    outcome = SimpleNamespace(
        document=GeminiDocument.model_validate_json(
            _answer(
                notify_party=_field("BETA LTD"),
                port_of_loading=_field(None),
                port_of_discharge=_field(None),
                gross_weight_kg=_field(None),
            )
        ),
        key_attempts=OK,
    )

    with pytest.raises(ExtractionFailure) as caught:
        grounded_extraction(document, outcome)
    assert caught.value.code is ExtractionFailureCode.UNGROUNDED_VALUE
    assert document.locate("BETA LTD", ComparedField.NOTIFY_PARTY) is None
    assert (
        _spot(document.locate("BETA LTD", ComparedField.CONSIGNEE))
        == spots["consignee"]
    )


def test_locate_treats_a_typographic_apostrophe_as_part_of_a_token():
    parsed = _parsed("Container Count: 2 x 40’HC\n".encode())

    assert parsed.locate("40") is None
    assert parsed.locate("40’HC").root.location.start_col == 21


def test_txt_continuation_lines_stay_under_their_label():
    parsed = _parsed(b"Shipper: ACME LTD\n  1 HARBOUR ROAD, SINGAPORE\nPOD: BUSAN\n")

    assert parsed.locate("SINGAPORE", ComparedField.PORT_OF_LOADING) is None
    assert parsed.locate("SINGAPORE", ComparedField.SHIPPER).root.location.line == 2
