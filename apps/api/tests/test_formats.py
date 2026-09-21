from hashlib import sha256
from io import BytesIO
from pathlib import Path

import pytest

from app.formats import detect_format, preflight

BUNDLE = Path(__file__).resolve().parents[3] / "data" / "sdoc-hackathon-bundle"
ATTACHMENTS = BUNDLE / "attachments"


def _read(name: str) -> bytes:
    return (ATTACHMENTS / name).read_bytes()


@pytest.mark.parametrize(
    ("name", "detected"),
    [
        ("email_001_SI.txt", "txt"),
        ("email_005_SI.xlsx", "xlsx"),
        ("email_291_BL.docx", "docx"),
        ("email_273_BL.pdf", "pdf"),
    ],
)
def test_preflight_hashes_sizes_and_detects_by_magic_bytes(name, detected):
    data = _read(name)

    # A misleading name must not change detection.
    result = preflight(data, file_name="renamed.bin")

    assert result.status == "OK"
    assert result.detected_format == detected
    assert result.content_hash == sha256(data).hexdigest()
    assert result.byte_size == len(data)
    assert result.scanned is False
    assert result.diagnostic is None


def test_scanned_pdf_is_readable_but_flagged_for_gemini():
    result = preflight(_read("email_512_SI.pdf"), file_name="email_512_SI.pdf")

    assert result.status == "OK"
    assert result.scanned is True
    assert result.page_count == 1


def _page_image(page, rect):
    """Place a small grey PNG over rect on a PyMuPDF page."""
    import pymupdf

    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 8, 8), False)
    pixmap.clear_with(200)
    page.insert_image(pymupdf.Rect(*rect), stream=pixmap.tobytes("png"))


def test_scanned_page_carrying_a_fax_header_line_is_a_scan():
    import pymupdf

    with pymupdf.open() as pdf:
        page = pdf.new_page()
        _page_image(page, (0, 40, 595, 842))
        page.insert_text((36, 24), "FAX FROM +65 6123 4567   21 SEP 2026 10:32   P.1/1")
        data = pdf.tobytes()

    result = preflight(data, file_name="fax.pdf")

    assert (result.status, result.scanned, result.page_count) == ("OK", True, 1)


def test_digital_cover_page_in_front_of_a_scanned_page_is_a_scan():
    import pymupdf

    cover_lines = (
        "TRANSMITTAL COVER SHEET",
        "To: Documentation Desk",
        "From: APRIL FINE PAPER TRADING PTE LTD",
        "Re: Shipping instruction for booking PSGSE4981829",
        "Pages: 2 including this cover",
    )
    with pymupdf.open() as pdf:
        cover = pdf.new_page()
        for index, text in enumerate(cover_lines):
            cover.insert_text((72, 72 + 20 * index), text)
        _page_image(pdf.new_page(), (0, 0, 595, 842))
        data = pdf.tobytes()

    result = preflight(data, file_name="cover.pdf")

    assert (result.status, result.scanned, result.page_count) == ("OK", True, 2)


def test_text_rich_page_with_a_small_logo_stays_digital():
    import pymupdf

    si_lines = (
        "BILL OF LADING INSTRUCTION",
        "B/L NUMBER: OOLU3584143842    BOOKING NO. PSGSE4981829",
        "Shipper",
        "APRIL FINE PAPER TRADING",
        "77 ROBINSON ROAD, #21-01, SINGAPORE 068896",
        "Consignee",
        "BALL & DOGGETT AUSTRALIA PTY LTD",
        "43-45 METROPOLITAN ROAD, ENFIELD NSW 2136, AUSTRALIA",
        "Notify Party",
        "PACIFIC OFFICE (M) SDN BHD",
        "POL",
        "BUATAN, INDONESIA",
        "Port of Discharge (POD)",
        "FREMANTLE, AUSTRALIA",
        "No. of Containers: 6 x 40'HC",
        "TOTAL Gross Wt (kgs): 131,322 KG",
    )
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        _page_image(page, (460, 36, 540, 76))  # the letterhead logo
        for index, text in enumerate(si_lines):
            page.insert_text((72, 100 + 18 * index), text)
        data = pdf.tobytes()

    result = preflight(data, file_name="si.pdf")

    assert (result.status, result.scanned) == ("OK", False)


@pytest.mark.parametrize("name", ["email_511_BL.pdf", "email_515_BL.pdf"])
def test_truncated_pdf_is_corrupt_with_a_diagnostic(name):
    result = preflight(_read(name), file_name=name)

    assert result.status == "CORRUPT"
    assert result.detected_format == "pdf"
    assert result.diagnostic.startswith("PDF could not be opened")


def test_damaged_ooxml_container_is_corrupt_not_unsupported():
    result = preflight(b"PK\x03\x04not-a-real-archive", file_name="draft.docx")

    assert result.status == "CORRUPT"
    assert result.diagnostic == "DOCX container is damaged"


def test_unknown_binary_is_unsupported():
    result = preflight(b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR", file_name="a.png")

    assert result.status == "UNSUPPORTED"
    assert result.detected_format == "unknown"
    assert result.diagnostic == "File type is not TXT, PDF, DOCX, or XLSX"


def test_detect_format_rejects_control_characters_in_text():
    assert detect_format(b"SHIPPING INSTRUCTION\x00") == "unknown"
    assert detect_format("毛重: 1 KG".encode()) == "txt"


from app.contracts import ComparedField
from app.formats import PreflightError, label_field, parse_document


def _parse(name: str):
    data = _read(name)
    return data, parse_document(
        data, preflight(data, file_name=name), attachment_id="att-1", file_name=name
    )


def _only(document, field):
    candidates = document.values()[field]
    assert len(candidates) == 1
    return candidates[0]


def test_txt_anchor_uses_code_points_for_chinese_labels():
    data, document = _parse("email_013_BL.txt")
    candidate = _only(document, ComparedField.GROSS_WEIGHT_KG)
    location = candidate.provenance.root.location
    line = data.decode("utf-8").split("\n")[location.line - 1]

    assert line == "Gross Weight毛重(KGS): 67,311 KG"
    assert (location.line, location.start_col, location.end_col) == (12, 21, 30)
    assert line[location.start_col : location.end_col] == "67,311 KG"
    # The UTF-8 byte offset is 25: a byte-offset bug would highlight the wrong text.
    assert len(line[: location.start_col].encode("utf-8")) == 25
    assert candidate.raw_value == "67,311 KG"


def test_txt_party_value_is_the_name_line_only():
    _, document = _parse("email_001_SI.txt")
    shipper = _only(document, ComparedField.SHIPPER)

    assert shipper.raw_value == "APRIL FAR EAST (M) SDN BHD"
    assert shipper.provenance.root.format == "txt"
    assert document.ambiguous_fields == ()


def test_blank_txt_value_keeps_a_zero_width_anchor_on_its_label_line():
    _, document = _parse("email_519_SI.txt")
    shipper = _only(document, ComparedField.SHIPPER)
    location = shipper.provenance.root.location

    assert shipper.raw_value == ""
    assert location.start_col == location.end_col


def _txt_document(text: str):
    data = text.encode()
    return parse_document(
        data, preflight(data, file_name="t.txt"), attachment_id="a", file_name="t.txt"
    )


def _txt_anchor(document, field):
    location = _only(document, field).provenance.root.location
    return location.line, location.start_col, location.end_col


def test_txt_line_numbers_count_only_newlines():
    # A viewer shows text.split("\n"): U+2028 and a form feed start no line.
    text = (
        "SHIPPING INSTRUCTION\n"
        "Page 1 of 2\u2028\n"
        "Shipper: ACME LTD\n"
        "\fConsignee: BETA LTD\n"
    )
    document = _txt_document(text)
    lines = text.split("\n")

    assert _txt_anchor(document, ComparedField.SHIPPER) == (3, 9, 17)
    assert _txt_anchor(document, ComparedField.CONSIGNEE) == (4, 12, 20)
    assert (lines[2][9:17], lines[3][12:20]) == ("ACME LTD", "BETA LTD")


def test_txt_crlf_file_anchors_like_its_lf_twin():
    text = "Shipper: ACME LTD\n  1 HARBOUR ROAD\n\n  SINGAPORE\nPOD: BUSAN\n"
    lf = _txt_document(text)
    crlf = _txt_document(text.replace("\n", "\r\n"))
    singapore = crlf.locate("SINGAPORE", ComparedField.PORT_OF_LOADING).root.location

    assert _txt_anchor(crlf, ComparedField.SHIPPER) == (1, 9, 17)
    assert _txt_anchor(crlf, ComparedField.PORT_OF_DISCHARGE) == (5, 5, 10)
    assert (singapore.line, singapore.start_col, singapore.end_col) == (4, 2, 11)
    assert [item.provenance for item in crlf.candidates] == [
        item.provenance for item in lf.candidates
    ]


def test_xlsx_anchor_is_sheet_and_value_cell():
    _, document = _parse("email_005_SI.xlsx")
    consignee = _only(document, ComparedField.CONSIGNEE)
    weight = _only(document, ComparedField.GROSS_WEIGHT_KG)

    assert consignee.raw_value == "BALL & DOGGETT AUSTRALIA PTY LTD"
    assert consignee.provenance.root.location.model_dump() == {
        "kind": "xlsx",
        "sheet": "S.I.",
        "cell": "B5",
    }
    assert weight.raw_value == "341715"
    assert weight.provenance.root.location.cell == "B10"


def test_docx_anchor_is_the_table_value_cell():
    _, document = _parse("email_291_BL.docx")
    shipper = _only(document, ComparedField.SHIPPER)
    weight = _only(document, ComparedField.GROSS_WEIGHT_KG)

    assert shipper.raw_value == "APRIL FINE PAPER TRADING (MIDDLE EAST) FZE"
    assert shipper.provenance.root.location.model_dump() == {
        "kind": "docx_table",
        "table_index": 0,
        "row_index": 0,
        "col_index": 1,
    }
    assert weight.raw_value == "21,745"
    assert weight.provenance.root.location.row_index == 6


def test_digital_pdf_anchor_is_page_and_value_bbox():
    _, document = _parse("email_273_BL.pdf")
    consignee = _only(document, ComparedField.CONSIGNEE)
    containers = _only(document, ComparedField.CONTAINER_COUNT)
    weight = _only(document, ComparedField.GROSS_WEIGHT_KG)
    location = consignee.provenance.root.location

    assert consignee.raw_value == "KPP-ANTALIS (SINGAPORE) PTE. LTD."
    assert location.page == 1
    assert location.approximate is False
    x0, y0, x1, y1 = location.bbox
    assert 160 < x0 < x1 < 330 and 120 < y0 < y1 < 150
    assert containers.raw_value == "2 x 20'FCL"
    assert weight.raw_value == "41,604 KG"
    assert document.ambiguous_fields == ()


@pytest.mark.parametrize(
    "name",
    [
        path.name
        for path in sorted(ATTACHMENTS.iterdir())
        if path.suffix != ".pdf"
        or path.name[6:9] not in {"511", "512", "513", "514", "515"}
    ],
)
def test_every_bundle_si_and_bl_resolves_all_seven_fields_locally(name):
    _, document = _parse(name)
    heading = document.text.strip().splitlines()[0].upper()
    if any(word in heading for word in ("INVOICE", "PACKING", "CERTIFICATE")):
        pytest.skip("not an SI or draft BL")

    assert document.ambiguous_fields == ()
    for candidate in document.candidates:
        location = candidate.provenance.root.location
        assert candidate.provenance.root.format in {
            "txt",
            "xlsx",
            "docx",
            "digital_pdf",
        }
        if location.kind == "txt":
            line = _read(name).decode("utf-8").split("\n")[location.line - 1]
            assert line[location.start_col : location.end_col] == candidate.raw_value


def test_missing_labels_make_a_local_parse_ambiguous():
    data = b"SHIPPING INSTRUCTION\nShipper: ACME LTD\nConsignee: BETA LTD\n"
    document = parse_document(
        data, preflight(data, file_name="x.txt"), attachment_id="a", file_name="x.txt"
    )

    assert ComparedField.SHIPPER not in document.ambiguous_fields
    assert ComparedField.NOTIFY_PARTY in document.ambiguous_fields


def test_conflicting_duplicate_labels_make_a_local_parse_ambiguous():
    data = b"Shipper: ACME LTD\nSHIPPER: OTHER LTD\n"
    document = parse_document(
        data, preflight(data, file_name="x.txt"), attachment_id="a", file_name="x.txt"
    )

    assert ComparedField.SHIPPER in document.ambiguous_fields


def test_locate_grounds_a_value_in_each_format():
    _, txt = _parse("email_001_BL.txt")
    _, pdf = _parse("email_273_BL.pdf")

    assert txt.locate("CALLAO, PERU (PECLL)").root.location.line == 10
    assert pdf.locate("CONAKRY, GUINEA").root.format == "digital_pdf"
    assert txt.locate("NOT IN THE DOCUMENT") is None
    assert txt.locate("   ") is None


def test_parse_document_refuses_scans_and_failed_preflight():
    scan = _read("email_512_SI.pdf")
    with pytest.raises(PreflightError):
        parse_document(
            scan,
            preflight(scan, file_name="s.pdf"),
            attachment_id="a",
            file_name="s.pdf",
        )


@pytest.mark.parametrize(
    ("label", "field"),
    [
        ("Shipper (Principal or Seller)", ComparedField.SHIPPER),
        ("To the Order of (收货人)", ComparedField.CONSIGNEE),
        ("Notify Party/Intermediate Consignee", ComparedField.NOTIFY_PARTY),
        ("Load Port (装货港)", ComparedField.PORT_OF_LOADING),
        ("POD", ComparedField.PORT_OF_DISCHARGE),
        ("No. of Containers or Packages", ComparedField.CONTAINER_COUNT),
        ("TOTAL Gross WeightII(KGS)", ComparedField.GROSS_WEIGHT_KG),
        ("Gross Wt (kgs) (毛重 KGS)", ComparedField.GROSS_WEIGHT_KG),
        ("NET WEIGHT", None),
        ("Vessel Name", None),
    ],
)
def test_label_field_aligns_labels_by_meaning(label, field):
    assert label_field(label) is field


def test_damaged_pdf_page_content_returns_corrupt_status(monkeypatch):
    """Test that PDF page content read exceptions are caught and return CORRUPT."""
    import pymupdf

    # Create a valid PDF that can be opened but fails on page.get_text()
    pdf_data = b"%PDF-1.4\n1 0 obj\n<</Type /Catalog /Pages 2 0 R>>\nendobj\n"
    pdf_data += b"2 0 obj\n<</Type /Pages /Kids [3 0 R] /Count 1>>\nendobj\n"
    pdf_data += (
        b"3 0 obj\n<</Type /Page /Parent 2 0 R /MediaBox [0 0 612 792]>>\nendobj\n"
    )
    pdf_data += b"xref\n0 4\n0000000000 65535 f\n0000000009 00000 n\n"
    pdf_data += b"0000000058 00000 n\n0000000115 00000 n\ntrailer\n"
    pdf_data += b"<</Size 4 /Root 1 0 R>>\nstartxref\n185\n%%EOF"

    def mock_get_text(*args, **kwargs):
        raise RuntimeError("Page content damaged")

    # Monkeypatch pymupdf.Page.get_text to raise an exception
    monkeypatch.setattr(pymupdf.Page, "get_text", mock_get_text)

    result = preflight(pdf_data, file_name="damaged.pdf")

    assert result.status == "CORRUPT"
    assert result.detected_format == "pdf"
    assert result.diagnostic.startswith("PDF could not be read")
    assert "RuntimeError" in result.diagnostic


def test_pdf_block_label_with_colon_extracts_correct_value():
    """Test that block labels followed by colons are parsed correctly."""
    import pymupdf

    # Create a small digital PDF with "NOTIFY PARTY: XYZ CO" on one line
    # and "Shipper" / "ACME LTD" on following lines
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "NOTIFY PARTY: XYZ CO")
    page.insert_text((72, 100), "Shipper")
    page.insert_text((72, 128), "ACME LTD")
    pdf_bytes = doc.tobytes()
    doc.close()

    pf = preflight(pdf_bytes, file_name="test.pdf")
    assert pf.status == "OK"

    document = parse_document(
        pdf_bytes, pf, attachment_id="att-1", file_name="test.pdf"
    )

    # Find the NOTIFY_PARTY candidate
    notify_party_candidates = document.values().get(ComparedField.NOTIFY_PARTY, [])
    assert len(notify_party_candidates) > 0, "NOTIFY_PARTY not found"

    notify_candidate = notify_party_candidates[0]
    # The raw_value should be just "XYZ CO", not " PARTY: XYZ CO"
    assert notify_candidate.raw_value == "XYZ CO"
    # The bbox x0 should be after the label (further right on the page)
    location = notify_candidate.provenance.root.location
    assert location.bbox[0] > 72  # x0 should be greater than label start

    # Also verify Shipper is parsed correctly from the next lines
    shipper_candidates = document.values().get(ComparedField.SHIPPER, [])
    assert len(shipper_candidates) > 0, "SHIPPER not found"
    assert shipper_candidates[0].raw_value == "ACME LTD"


def _pdf_document(*lines: str):
    """A one-page digital PDF with each text on its own baseline, 28pt apart."""
    import pymupdf

    with pymupdf.open() as pdf:
        page = pdf.new_page()
        for index, text in enumerate(lines):
            page.insert_text((72, 72 + 28 * index), text)
        data = pdf.tobytes()
    return parse_document(
        data, preflight(data, file_name="t.pdf"), attachment_id="a", file_name="t.pdf"
    )


def test_blank_pdf_block_label_does_not_take_a_label_line_as_its_value():
    document = _pdf_document("Notify Party", "Gross Weight: 12,000 KG")
    notify = _only(document, ComparedField.NOTIFY_PARTY)
    _, y0, _, y1 = notify.provenance.root.location.bbox

    assert notify.raw_value == ""
    assert y0 < 72 < y1  # the blank is anchored on its own label line
    assert _only(document, ComparedField.GROSS_WEIGHT_KG).raw_value == "12,000 KG"


@pytest.mark.parametrize(
    "header",
    [
        "Vessel",
        "Vessel Name",
        "Ocean Vessel",
        "Export Carrier (vessel, voyage)SOLID 16 V.044NW2",
        "CONTAINER NO.",
        "DESCRIPTION",
        "GROSS WEIGHT (KG)",
        "HS CODE 48025700   FREIGHT PREPAID",
        "B/L NUMBER",
        "BOOKING NO. PSGSE4981829",
        "Place of Receipt",
    ],
)
def test_blank_pdf_block_label_does_not_take_a_section_header_as_its_value(header):
    document = _pdf_document("Port of Loading", header, "SINGAPORE")

    assert _only(document, ComparedField.PORT_OF_LOADING).raw_value == ""


def _docx_document(source):
    """Parse a python-docx document built in memory."""
    buffer = BytesIO()
    source.save(buffer)
    data = buffer.getvalue()
    return parse_document(
        data, preflight(data, file_name="t.docx"), attachment_id="a", file_name="t.docx"
    )


def test_docx_label_merged_across_columns_takes_the_next_distinct_cell():
    import docx

    source = docx.Document()
    table = source.add_table(rows=1, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Shipper"
    table.cell(0, 2).text = "ACME LTD"

    shipper = _only(_docx_document(source), ComparedField.SHIPPER)

    assert shipper.raw_value == "ACME LTD"
    assert shipper.provenance.root.location.model_dump() == {
        "kind": "docx_table",
        "table_index": 0,
        "row_index": 0,
        "col_index": 2,
    }


def test_docx_label_merged_across_the_whole_row_is_a_blank_value():
    import docx

    source = docx.Document()
    table = source.add_table(rows=1, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "Notify Party"

    notify = _only(_docx_document(source), ComparedField.NOTIFY_PARTY)

    assert notify.raw_value == ""
    assert notify.provenance.root.location.col_index == 0  # anchored on the label


def _si_workbook():
    """Shipper merged across A1:B1 with its value in C1, a gross weight formula
    saved (as openpyxl saves it) with no cached result, and a blank consignee."""
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "Shipper"
    sheet.merge_cells("A1:B1")
    sheet["C1"] = "ACME TRADING"
    sheet["A2"] = "Gross Weight"
    sheet["B2"] = "=10+5"
    sheet["A3"] = "Consignee"
    buffer = BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()
    return parse_document(
        data, preflight(data, file_name="t.xlsx"), attachment_id="a", file_name="t.xlsx"
    )


def test_xlsx_label_merged_across_columns_takes_the_next_cell_outside_it():
    document = _si_workbook()
    shipper = _only(document, ComparedField.SHIPPER)

    assert (shipper.raw_value, shipper.provenance.root.location.cell) == (
        "ACME TRADING",
        "C1",
    )
    assert ComparedField.SHIPPER not in document.ambiguous_fields


def test_xlsx_formula_without_a_cached_result_is_ambiguous_not_blank():
    document = _si_workbook()

    assert ComparedField.GROSS_WEIGHT_KG in document.ambiguous_fields
    assert ComparedField.GROSS_WEIGHT_KG not in document.values()


def test_xlsx_formula_without_a_cached_result_unsettles_a_value_read_elsewhere():
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.active["A1"], workbook.active["B1"] = "Gross Weight", "=10+5"
    packing = workbook.create_sheet("Packing")
    packing["A1"], packing["B1"] = "Gross Weight", 15
    buffer = BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()
    document = parse_document(
        data, preflight(data, file_name="t.xlsx"), attachment_id="a", file_name="t.xlsx"
    )

    assert ComparedField.GROSS_WEIGHT_KG in document.ambiguous_fields


def test_xlsx_empty_value_cell_is_still_a_settled_blank():
    document = _si_workbook()
    consignee = _only(document, ComparedField.CONSIGNEE)

    assert (consignee.raw_value, consignee.provenance.root.location.cell) == ("", "B3")
    assert ComparedField.CONSIGNEE not in document.ambiguous_fields
